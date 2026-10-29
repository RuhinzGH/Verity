"""Plain text out of uploaded files: .txt, .md, .pdf and .docx.

Only the text is read. Images, scanned pages (a PDF that is a picture of text) and the
layout of the page are not, so a scanned PDF comes back empty.
"""
from __future__ import annotations

import io
import re

TYPES = ["txt", "md", "pdf", "docx"]


def read_text(name: str, data: bytes) -> str:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext == "pdf":
        return _pdf(data)
    if ext == "docx":
        return _docx(data)
    return data.decode("utf-8", errors="replace").replace("\r\n", "\n").strip()


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    return unwrap("\n\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages))


def _docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:                     # each row becomes one line of text
        for row in table.rows:
            cells = []
            for c in row.cells:                  # merged cells repeat; keep each once
                if not cells or cells[-1] != c.text:
                    cells.append(c.text)
            lines.append("; ".join(t.strip() for t in cells if t.strip()))
    return "\n".join(l.strip() for l in lines if l.strip())


def unwrap(text: str) -> str:
    """Join the lines a PDF breaks mid-sentence: Verity treats every newline as the end of
    a sentence. Blank lines (paragraph breaks) are kept as one newline; a word hyphenated
    across a line end is joined."""
    paras = re.split(r"\n\s*\n", text.replace("\r\n", "\n"))
    out = []
    for p in paras:
        p = re.sub(r"(\w)-\n([a-z])", r"\1\2", p)
        p = re.sub(r"\s*\n\s*", " ", p)
        p = re.sub(r"[ \t]{2,}", " ", p).strip()
        if p:
            out.append(p)
    return "\n".join(out)
