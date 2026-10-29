"""Text extraction from uploaded files (offline: the files are built in memory)."""
import io

from verity.files import read_text, unwrap


def _pdf(lines: list[str]) -> bytes:
    """A minimal one-page PDF with one text line per entry (Helvetica)."""
    text = " ".join(f"({l}) Tj 0 -16 Td" for l in lines)
    stream = f"BT /F1 12 Tf 72 720 Td {text} ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + o + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_txt_and_md_are_decoded():
    assert read_text("a.txt", "Water boils.\r\nIt is wet.".encode()) == "Water boils.\nIt is wet."
    assert read_text("notes.MD", b"# Title") == "# Title"


def test_pdf_lines_are_joined_into_sentences():
    data = _pdf(["Pure water boils at 100 degrees", "Celsius. It freezes at 0 degrees Celsius."])
    assert read_text("water.pdf", data) == \
        "Pure water boils at 100 degrees Celsius. It freezes at 0 degrees Celsius."


def test_docx_paragraphs_and_tables():
    from docx import Document
    d = Document()
    d.add_paragraph("Mars is the fourth planet from the Sun.")
    d.add_paragraph("")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Moons", "Phobos and Deimos"
    buf = io.BytesIO(); d.save(buf)
    assert read_text("mars.docx", buf.getvalue()) == \
        "Mars is the fourth planet from the Sun.\nMoons; Phobos and Deimos"


def test_unwrap_keeps_paragraphs_and_joins_hyphenated_words():
    assert unwrap("A long sen-\ntence here.\n\n\nNext para-\ngraph\nends.") == \
        "A long sentence here.\nNext paragraph ends."
